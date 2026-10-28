"""Build private, review-first delivery reports without changing academic records.

Input packages come from archive_release. This step does not grant rights, approve
questions, infer examination targets, or change any WordPress database.
"""
import argparse
import csv
import gzip
import html
import json
import re
import shutil
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
import archive_catalog as catalog
import fr_package

ROOT = Path(__file__).resolve().parents[1]
SUBJECTS = ('FR', 'AFM', 'AUDIT', 'DT', 'IDT', 'IBS')
COLOURS = dict(FR='E8F3FC', AFM='ECF6ED', AUDIT='F1EDFA', DT='FFF4E4', IDT='E8F7F4', IBS='FCEEF3')
STANDARDS = json.loads((ROOT/'ca-final-command-center/data/standards-reference.json').read_text())['entries']


def tags(row):
    """Explicit source references only, separately labelled from official headings."""
    r = row['import_record']
    question = row['question']['plain_text']
    answer = row['answer']['plain_text'] if row['answer'] else ''
    heading = r.get('heading_evidence', {}).get('text', '')
    proposals = []
    for standard in STANDARDS:
        if r['subject_code'] not in (standard['subject'], 'IBS'):
            continue
        pattern = (r'\bInd\s*AS\s*' if standard['family'] == 'Ind AS' else r'\bSA\s*') + str(standard['number']) + r'\b'
        evidence = [side for side, text in [('heading', heading), ('question', question), ('candidate answer', answer)] if re.search(pattern, text, re.I)]
        if evidence:
            proposals.append(dict(label=standard['code'], name=standard['title'], origin='explicit_source_reference_provisional_applicability', evidence=evidence, naming_reference=standard['source_url']))
    for label in r['ai_proposal'].get('proposed_search_tags', []):
        proposals.append(dict(label=label, origin='source_heading_or_scoped_correction_proposal', evidence=['Exact private heading/correction evidence in question record']))
    if r['subject_code'] in ('DT', 'IDT', 'IBS'):
        for match in re.finditer(r'\b(?:section|rule|article)\s+\d+[A-Za-z]*(?:\s*\([A-Za-z0-9]+\))*', question, re.I):
            proposals.append(dict(label=match[0], origin='printed_statutory_reference_not_verified_primary_topic', evidence=['Question text; historical law remains unchanged']))
    return dict(stable_id=r['stable_id'], subject=r['subject_code'], scheme=r['scheme'], attempt=r['exam_attempt'], law_context=r['law_context'], proposals=proposals, official_historical_mapping='Not independently reconciled', fallback=r['ai_proposal']['primary_path'][0]['title'])


def sheet(book, name, headers, rows, colour):
    if name in book:
        del book[name]
    page = book.create_sheet(name)
    page.append(headers)
    for row in rows:
        page.append([str(x) if isinstance(x, (dict, list)) else x for x in row])
    page.freeze_panes = 'A2'
    page.auto_filter.ref = page.dimensions
    for cell in page[1]:
        cell.fill = PatternFill('solid', fgColor='102542')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in page.iter_rows(min_row=2):
        for cell in row:
            if cell.data_type == 'f':
                cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.fill = PatternFill('solid', fgColor=colour)
    for col in page.columns:
        page.column_dimensions[col[0].column_letter].width = 32
    return page


def excel_text(value, stable_id):
    return value if len(value) <= 32767 else 'Full text exceeds the Excel cell limit. Open private details/'+stable_id+'.html. Complete text is retained; no truncation.'


def build(manifest_path, worker_path, release_path, output, revision):
    release = catalog.private(release_path)
    out = catalog.private(output)
    out.mkdir(parents=True, exist_ok=False)
    manifest = json.loads(Path(manifest_path).read_text())
    workers = json.loads(Path(worker_path).read_text())
    by_id = {r['file_id']: r for r in workers['documents']}
    summaries = json.loads((release/'processing-summary.json').read_text())
    # Exclude machine-specific paths from portable register; exact file IDs and
    # source hashes remain sufficient to retrieve the original authorised PDF.
    portable = []
    for original in manifest['records']:
        r = {k: v for k, v in original.items() if k not in ('local_path', 'evidence_path')}
        r['outcome'] = by_id.get(r['file_id'], {}).get('status', 'not_extracted')
        r['student_pdf_url'] = 'https://drive.google.com/file/d/'+r['file_id']+'/view'
        r['rights_scope'] = 'No new grant. Match the actual administrator-recorded permission before publication.'
        r['source_evidence'] = 'evidence/'+r['source_hash']+'/pages.json.gz'
        portable.append(r)
    catalog.write(out/'master-source-register.json', dict(records=portable, counts=manifest.get('counts', {}), worker_counts=workers['counts'], pending=workers.get('pending', 0)))
    fields = ['file_id', 'filename', 'collection_subject', 'inventory_group', 'bytes', 'source_hash', 'page_count', 'actual_subject', 'source_type', 'role', 'exam_attempt', 'document_issue_period', 'series', 'scheme', 'assessment_year_evidence', 'outcome', 'student_pdf_url', 'rights_scope', 'warnings']
    flat = []
    for r in portable:
        d = {**r, **r['identity']}
        flat.append({k: '; '.join(map(str, d[k])) if isinstance(d.get(k), list) else d.get(k, '') for k in fields})
    with (out/'master-source-register.csv').open('w', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader(); writer.writerows(flat)
    catalogue_book = Workbook(); catalogue_book.remove(catalogue_book.active)
    sheet(catalogue_book, 'Source Register', fields, [[r[k] for k in fields] for r in flat], 'EDF3F8')
    catalogue_book.save(out/('Master-Source-Register-'+revision+'.xlsx'))
    delivery = []
    for subject in SUBJECTS:
        base = release/subject
        rows = json.loads((base/'subjective-candidates.json').read_text())
        mcqs = json.loads((base/'mcq-candidates.json').read_text())
        for c in mcqs:
            options = ''.join('<li>'+o['html']+'</li>' for o in c['options'])
            (base/'details'/(c['stable_id']+'.html')).write_text('<!doctype html><meta charset="utf-8"><h1>PRIVATE MCQ — Needs Review</h1><p>'+html.escape(c['stable_id'])+'</p>'+((c.get('scenario') or {}).get('html', ''))+c['question_html']+'<ol type="A">'+options+'</ol><h2>Provisional key/answer evidence</h2><pre>'+html.escape(json.dumps(c['answer'], ensure_ascii=False, indent=2))+'</pre>'+c.get('explanation_html', ''))
        tag_report = [tags(row) for row in rows]
        catalog.write(base/'tagging-and-historical-report.json', tag_report)
        book_path = next(base.glob('*Owner-Review-*.xlsx'))
        book = load_workbook(book_path)
        colour = COLOURS[subject]
        sheet(book, 'Source Inventory', fields, [[r[k] for k in fields] for r in flat if r['collection_subject'] == subject], colour)
        for i, row in enumerate(rows, 2):
            r = row['import_record']
            book['Questions'].cell(i, 17, r['question_type'])
            book['Questions'].cell(i, 13, '; '.join(t['label'] for t in tag_report[i-2]['proposals']))
            for cell in book['Questions'][i]:
                cell.fill = PatternFill('solid', fgColor=colour)
        sheet(book, 'Mapping Evidence', ['Stable ID', 'Historical scheme', 'Historical law', 'Proposed tags', 'Provenance', 'Import identity'], [[r['import_record']['stable_id'], r['import_record']['scheme'], r['import_record']['law_context'], '; '.join(t['label'] for t in tgs['proposals']), json.dumps(tgs['proposals'], ensure_ascii=False), 'Importable draft' if r['import_eligible'] else 'HELD: source identity unresolved'] for r, tgs in zip(rows, tag_report)], colour)
        sheet(book, 'MCQ Review', ['Stable ID', 'Source', 'Attempt', 'Series', 'Reference', 'Full stem / shared facts', 'Exact options', 'Candidate key', 'Key evidence', 'Warnings', 'Review decision', 'Owner corrections'], [[c['stable_id'], c['source_type'], c['exam_attempt'], c['mtp_series'], c['reference'], excel_text(html.unescape(re.sub('<[^>]+>', ' ', (c.get('scenario') or {}).get('html', '')+c['question_html'])), c['stable_id']), excel_text(json.dumps(c['options'], ensure_ascii=False), c['stable_id']), ', '.join(c['correct_options']), excel_text(json.dumps(c['answer'].get('key_evidence'), ensure_ascii=False), c['stable_id']), '; '.join(c['exceptions']), 'Needs Review', ''] for c in mcqs], colour)
        if mcqs:
            validation = DataValidation(type='list', formula1='"Needs Review,Publish,Reject"', allow_blank=True)
            book['MCQ Review'].add_data_validation(validation)
            validation.add('K2:K'+str(len(mcqs)+1))
        summary = summaries[subject]
        sheet(book, 'Processing Coverage', ['Measure', 'Value'], [[k, v] for k, v in summary.items() if k != 'packages'], colour)
        instructions = [
            ['Private research, not academic approval', 'All records remain Needs Review. Source context/heading tags are provisional. Source-copy access does not extend publication rights.'],
            ['Import', 'Extract the subject outer ZIP. Upload imports/*.zip separately through Upload Questions. Each child contains its own taxonomy/sources in dependency order. Save as Drafts first.'],
            ['Existing January pilot', 'The seven reserved January FR PDF identities and content-identical copies are excluded from new imports. Retain the existing 55 subjective / 54 eligible MCQ records.'],
            ['Excel decisions', 'This main workbook is the complete research overview. Use decisions/*.xlsx (at most 150 identity-eligible subjective records each) for the existing WordPress validation/upload flow. Import matching drafts first. Existing or changed records require a freshly exported WordPress workbook; stale uploads are rejected. No approval is prefilled. Never upload the larger overview workbook as a decision batch.'],
            ['MCQ decisions', 'MCQ Review is private research/annotation. Publish eligible MCQs through the owner package workflow; these extra MCQ sheet decisions are not silently applied by the subjective Excel importer.'],
            ['Long records', 'Complete structured HTML is in details/<Stable ID>.html and JSON. Do not publish a record with an unresolved material table/boundary conflict.'],
            ['Held sources', 'Missing target attempt, series, historical subject mapping, scanned pages and missing keys require targeted evidence. Do not substitute the issue date or latest syllabus.'],
            ['Rights', 'The import supplies no permission grant. Match the actual source-wide physical-authorisation attestation; no scan is mandatory.'],
        ]
        sheet(book, 'Read First', ['Action', 'Instructions'], instructions, colour)
        book.save(book_path)
        shutil.copy2(book_path, out/book_path.name)
        decision_dir = base/'decisions'
        decision_dir.mkdir(exist_ok=True)
        import_rows = [r for r in rows if r['import_eligible']]
        for first in range(0, len(import_rows), 150):
            batch = import_rows[first:first+150]
            name = fr_package.workbook(batch, [], {'nodes': []}, [], decision_dir, first//150+1)
            (decision_dir/name).rename(decision_dir/(subject+'-Revision-Safe-Decisions-'+str(first//150+1).zfill(3)+'.xlsx'))
        # Import-only categories are independent child archives. The outer ZIP is
        # a delivery container, deliberately not a nested WordPress import.
        catalog.write(base/'private-review-status.json', dict(revision=revision, academic_approvals=0, automatic_publications=0, complete_corpus=False, software_import_validation='See private-package-validation.json', table_validation='Automated numeric/geometry checks only; independent visual checks are explicitly listed, never assumed.'))
        (base/'READ-FIRST.txt').write_text('\n\n'.join(title+'\n'+text for title, text in instructions)+'\n\nRead processing-summary.json, source-pairing-report.json, table-report.json, exceptions.json and import-order.json. Processing counts are not publication counts.\n')
        checks = []
        for path in sorted(base.rglob('*')):
            if path.is_file() and path.name != 'SHA256SUMS.txt':
                checks.append(catalog.sha(path.read_bytes())+'  '+str(path.relative_to(base)))
        (base/'SHA256SUMS.txt').write_text('\n'.join(checks)+'\n')
        target = out/(subject+'-ALL-ATTEMPTS-v0.8.1-'+revision+'-PRIVATE-NEEDS-REVIEW.zip')
        with zipfile.ZipFile(target, 'x', zipfile.ZIP_DEFLATED) as z:
            for path in sorted(base.rglob('*')):
                if path.is_file(): z.write(path, str(path.relative_to(base)))
            evidence_checks = []
            for r in manifest['records']:
                if r['collection_subject'] == subject and r['file_id'] in by_id:
                    path = Path(by_id[r['file_id']]['payload_path'])
                    name = 'layout-evidence/'+r['file_id']+'.json'
                    z.write(path, name)
                    evidence_checks.append(catalog.sha(path.read_bytes())+'  '+name)
            z.writestr('LAYOUT-SHA256SUMS.txt', '\n'.join(evidence_checks)+'\n')
        for path in (target, out/book_path.name):
            delivery.append(dict(filename=path.name, bytes=path.stat().st_size, sha256=catalog.sha(path.read_bytes()), subject=subject, academic_approvals=0))
    evidence_zip = out/('Six-Subject-Source-Evidence-'+revision+'-PRIVATE.zip')
    with zipfile.ZipFile(evidence_zip, 'x', zipfile.ZIP_DEFLATED) as z:
        for name in ('master-source-register.json', 'master-source-register.csv', 'Master-Source-Register-'+revision+'.xlsx'):
            z.write(out/name, name)
        seen = set()
        for r in manifest['records']:
            p = Path(r['evidence_path'])
            if r['source_hash'] not in seen:
                z.write(p, 'evidence/'+r['source_hash']+'/pages.json.gz')
                seen.add(r['source_hash'])
        z.writestr('coverage.json', json.dumps(summaries, ensure_ascii=False, indent=2))
        z.writestr('worker-outcomes.json', json.dumps({k:v for k,v in workers.items() if k!='documents'}, indent=2))
    delivery.append(dict(filename=evidence_zip.name, bytes=evidence_zip.stat().st_size, sha256=catalog.sha(evidence_zip.read_bytes())))
    catalog.write(out/'delivery-manifest.json', dict(revision=revision, files=delivery, subjects=summaries, automatic_publications=0, academic_approvals=0, completion='Draft expansion with explicit held sources and unresolved academic/layout exceptions; not a fully verified six-subject bank.'))
    (out/'SHA256SUMS.txt').write_text('\n'.join(d['sha256']+'  '+d['filename'] for d in delivery)+'\n')
    return delivery


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('manifest'); p.add_argument('worker_manifest'); p.add_argument('release'); p.add_argument('output'); p.add_argument('--revision', required=True)
    args = p.parse_args()
    print(json.dumps(build(args.manifest, args.worker_manifest, args.release, args.output, args.revision), indent=2))
