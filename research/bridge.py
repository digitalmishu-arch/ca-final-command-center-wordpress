"""Metadata-only draft transport and private XLSX review books for the v0.4 Studio.

No question/answer prose or Drive links are copied into the WordPress transport.
Offline workbooks are research aids; server-issued workbooks alone can apply approvals.
"""
import argparse
from collections import Counter
from datetime import datetime,timezone
import hashlib,json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill,Protection
from openpyxl.worksheet.datavalidation import DataValidation

ROOT=Path(__file__).resolve().parents[1]
COLORS={'FR':'087F8C','AFM':'7152C8','AUDIT':'B86525','DT':'338254','IDT':'376FC3','IBS':'B64F83'}
SHEETS=['Question Review','Answer Review','Syllabus Mapping','Source Register','Exceptions','Approval Decisions','Batch Manifest','Progress Summary']

def packet(record):
    sid='R-'+record['question_identity'];c=record['classification'];candidates=record['answer_candidates'];answer={}
    if len(candidates)==1:
        a=candidates[0];answer=dict(source_hash=a['source_sha256'],official_url=a['official_url'],reference=record['reference'],subreference=record['subreference'],page=str(a['page']),coordinates=dict(page=a['page'],bbox=a['bbox']),question_identity=sid,question_reference=record['reference'],question_subreference=record['subreference'],notes='Structural candidate only. Owner must inspect the original answer and approve its exact location.')
    reasoning=dict(proposed_concept=c.get('concept'),method=c.get('method'),reasoning=c.get('reasoning'),official_nodes=[dict(key=n['node_key'],title=n['title'],source_url=n['source_url'],edition=n['edition_key']) for n in c.get('official_nodes',[])],editorial_topics=c.get('editorial_topics',[]),editorial_subtopics=c.get('editorial_subtopics',[]),additional_supported_concepts=c.get('additional_supported_concepts',[]),historical_edition=c.get('historical_edition'))
    return dict(stable_id=sid,source_hash=record['source_sha256'],source_key=record['document_key'],subject_code=record['subject'],source_type=record['source_type'],exam_attempt=record['attempt'],mtp_series=record.get('mtp_series',''),section=record['section'],reference=record['reference'],subreference=record['subreference'],title=record['source_title'],official_url=record['official_url'],scheme='',edition_id=None,law_context='',question_page=str(record['start']['page']),question_coordinates=dict(start=record['start'],end_exclusive=record['end_exclusive']),heading_evidence=record.get('heading') or {},classification=[],reasoning=json.dumps(reasoning,ensure_ascii=False),ai_proposal=reasoning,confidence='Provisional; not academic verification',exceptions=record['warnings'],answer=answer,review_notes='Historical scheme/edition and proposed tags must be assigned/reviewed in WordPress. Private source PDFs remain outside the website.',rights_mode='reference_only',rights_basis='Reference metadata only. No grant of ICAI question/answer reproduction rights is asserted.',extraction_status='needs_review',engine_revision=record.get('engine_revision',{}),previous_corrections=record.get('ledger_event') or {},marks=record.get('marks') if record.get('marks_verified') else None,marks_verified=False,question_type=record['kind'] if record['kind'] in ['practical','theory','mcq','case_study','case_scenario','unknown'] else 'unknown')

def export(review_path,inventory_path,output,batch_size=100):
    out=Path(output).resolve()
    if (ROOT/'work').resolve() not in out.parents or out.exists():raise ValueError('Use a NEW ignored private output directory.')
    if not 1<=batch_size<=100:raise ValueError('Batch size must be 1–100.')
    out.mkdir(parents=True,mode=0o700);records=json.loads(Path(review_path).read_text());inventory=json.loads(Path(inventory_path).read_text());packets=[packet(r) for r in records]
    for start in range(0,len(packets),batch_size):
        batch=dict(batch_key='RESEARCH-'+hashlib.sha256(Path(review_path).read_bytes()).hexdigest()[:16]+'-'+str(start//batch_size+1),records=packets[start:start+batch_size]);(out/f'studio-drafts-{start//batch_size+1:03d}.json').write_text(json.dumps(batch,indent=2,ensure_ascii=False))
    wb=Workbook();wb.remove(wb.active)
    for name in SHEETS:wb.create_sheet(name)
    q=wb['Question Review'];q.append(['Stable ID','Source SHA-256','Subject','Attempt','RTP/MTP/Exam','Series','Source PDF','Question Number','Subpart','Full permitted Question Text','Full permitted Answer Text','Chapter','Unit','Topic','Subtopic','Other Tags','Official URL','Question Page','Answer Page','AI Classification','Correction','Approve / Correct / Reject / Hold','Reviewer Notes'])
    a=wb['Answer Review'];a.append(['Stable ID','Answer Candidates','Verification','Owner Notes'])
    sm=wb['Syllabus Mapping'];sm.append(['Stable ID','Historical Edition Proposal','Official Chapter/Unit Proposals','Editorial Topics','Editorial Subtopics','Owner Decision'])
    sr=wb['Source Register'];sr.append(['Inventory File ID','Subject','Kind','Filename','Attempt Label','Scheme (unverified)','Availability','Private Research Rights','Public Reproduction Rights','Processing'])
    ex=wb['Exceptions'];ex.append(['Stable ID / File ID','Exception','State'])
    ad=wb['Approval Decisions'];ad.append(['Stable ID','Source Hash','Decision','Correction','Reviewer Notes','WordPress Status'])
    bm=wb['Batch Manifest'];bm.append(['Key','Value']);bm.append(['Format','CAFCC offline research review — NOT a server-issued approval workbook']);bm.append(['Created',datetime.now(timezone.utc).isoformat()]);bm.append(['Review source SHA-256',hashlib.sha256(Path(review_path).read_bytes()).hexdigest()]);bm.append(['Instructions','Keep private. No copied PDF prose is included. Review candidates against your authorised PDFs. Import studio-drafts JSON through administrator preview first. Then export a fresh server-issued Review Excel to apply decisions securely; this offline workbook cannot bypass server identity/revision checks.']);bm.append(['Publication','All records are Draft/Needs Review; no owner approvals created.']);bm.append(['AFM conflict','Raw Q1–Q4 Securitisation corrections remain unapplied pending exact section/subpart identification.'])
    for r,p in zip(records,packets):
        c=r['classification'];nodes=c.get('official_nodes',[]);chapter='; '.join(n['title'] for n in nodes if n['node_type']=='chapter');unit='; '.join(n['title'] for n in nodes if n['node_type']=='unit');topics='; '.join(c.get('editorial_topics',[]));subs='; '.join(c.get('editorial_subtopics',[]));aps='; '.join(str(a['page']) for a in r['answer_candidates'])
        q.append([p['stable_id'],p['source_hash'],r['subject'],r['attempt'],r['source_type'],r.get('mtp_series',''),r['document_key'],r['reference'],r['subreference'],'Reference-only; inspect authorised source PDF','No copied answer text',chapter,unit,topics,subs,'; '.join(c.get('additional_supported_concepts',[])),r['official_url'],r['start']['page'],aps,json.dumps(c,ensure_ascii=False),'','Hold',''])
        for cell in q[q.max_row]:cell.fill=PatternFill('solid',fgColor='F5F8FC')
        q.cell(q.max_row,3).fill=PatternFill('solid',fgColor=COLORS[r['subject']]);q.cell(q.max_row,3).font=Font(color='FFFFFF',bold=True)
        a.append([p['stable_id'],json.dumps(r['answer_candidates'],ensure_ascii=False),'Not owner verified','']);sm.append([p['stable_id'],c.get('historical_edition',''),chapter+' | '+unit,topics,subs,'Needs Review']);ad.append([p['stable_id'],p['source_hash'],'Hold','','','Not imported; offline decisions are provisional'])
        for warning in r['warnings']:ex.append([p['stable_id'],json.dumps(warning,ensure_ascii=False),'Needs Review'])
    source_keys={r['document_key'] for r in records}
    for r in inventory['records']:sr.append([r['file_id'],r['subject'],r['document_kind'],r['filename'],r['attempt'],r['scheme'],r['availability'],'Unconfirmed except existing six-file authorisation','Not established','See inventory coverage report; listing does not mean processed'])
    for row in inventory['pairing']:
        if row['status']!='candidate_only_verify_visible_attempt_series':ex.append([row['file_id'],row['status'],'Needs Review'])
    ps=wb['Progress Summary'];ps.append(['Measure','Value']);ps.append(['Inventoried workbook files',inventory['catalogued_rows']]);ps.append(['Matching accessible file listings',inventory['availability'].get('listed_downloadable',0)]);ps.append(['Processed source PDFs',len(source_keys)]);ps.append(['Provisional question/subpart records',len(records)]);ps.append(['Owner approved',0]);ps.append(['Published',0])
    for subject,count in Counter(r['subject'] for r in records).items():ps.append([subject+' provisional records',count])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False;ws.protection.sheet=True
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='102542');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[1].height=38
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value is not None and len(str(cell.value))>32767:raise ValueError('Excel cell exceeds its limit; supply a secure private detail reference instead of truncating.')
                if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
            ws.row_dimensions[row[0].row].height=78
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=30
    for row in q.iter_rows(min_row=2):
        for index in (21,22,23):row[index-1].protection=Protection(locked=False)
    for row in ad.iter_rows(min_row=2):
        for index in (3,4,5):row[index-1].protection=Protection(locked=False)
    decision=DataValidation(type='list',formula1='"Approve,Correct,Reject,Hold"');q.add_data_validation(decision);decision.add('V2:V'+str(q.max_row));decision2=DataValidation(type='list',formula1='"Approve,Correct,Reject,Hold"');ad.add_data_validation(decision2);decision2.add('C2:C'+str(ad.max_row))
    for col in ('J','K','T','U','W'):q.column_dimensions[col].width=65
    wb.save(out/'CA-Final-Private-Research-Review.xlsx')
    summary=dict(records=len(records),subjects=dict(Counter(r['subject'] for r in records)),transport_files=(len(packets)+batch_size-1)//batch_size,rights='reference_only',owner_approved=0,published=0)
    (out/'bridge-report.json').write_text(json.dumps(summary,indent=2));return summary

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--review',required=True);p.add_argument('--inventory',required=True);p.add_argument('--output',required=True);p.add_argument('--batch-size',type=int,default=100);a=p.parse_args();print(json.dumps(export(a.review,a.inventory,a.output,a.batch_size)))
