"""Build private FR owner-review data, rich details, small imports and XLSX books.

No academic approval, exact-repeat label or public source prose is produced.
Question counts are candidates, with parents and subparts reported separately.
"""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import html
import json
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

import fr_corpus

ROOT=Path(__file__).resolve().parents[1]


def sha(value):return hashlib.sha256(value if isinstance(value,bytes) else value.encode()).hexdigest()
def write(path,value):path.parent.mkdir(parents=True,exist_ok=True);path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n')


def match_answers(document,question,documents):
    identity=document['identity'];pool=[]
    if identity['role']=='questions_and_answers':pool=[document]
    elif identity['source_type']=='mtp':
        # Filename pairs are only inventory hints. Actual issue/series/subject
        # evidence is required independently on BOTH PDFs before candidate pairing.
        pool=[d for d in documents if d['identity']['role']=='answers' and d['identity']['source_type']=='mtp'
              and identity['subject_confirmed'] and d['identity']['subject_confirmed']
              and identity['issue_period'] and d['identity']['issue_period']==identity['issue_period']
              and d['identity']['series']==identity['series']]
    candidates=[]
    for doc in pool:
        for answer in doc['records']:
            if answer['side']=='answer' and not(question.get('format')=='SUBJECTIVE' and answer.get('format')=='MCQ') and answer['reference']==question['reference'] and answer['subreference']==question['subreference']:
                candidates.append(dict(document=doc,record=answer))
    return candidates


def classification(document,question,taxonomy,previous):
    ident=document['identity'];attempt=ident['attempt'];year=int(attempt[-4:]) if attempt else 0
    edition='fr-may2025-jan2026' if attempt in ('May 2025','September 2025','January 2026') else 'fr-may2026' if year==2026 and attempt!='January 2026' else ''
    # A heading supports a concept, never an invented historical chapter number.
    heading=question['heading'];standard=re.search(r'Ind(?:ian Accounting Standard)?\s*AS?\s*(\d+)',heading['text'] if heading else '',re.I)
    if not standard:standard=re.search(r'Ind\s+AS\s*(\d+)',heading['text'] if heading else '',re.I)
    concept='Ind AS '+standard[1] if standard else ''
    old=previous.get(question['stable_id']);ledger=old.get('ledger_event') if old else None
    if ledger and ledger.get('decision') in ('approved','concept_approved'):
        concept=ledger.get('classification',{}).get('concept',concept)
    official=[];detail=[]
    if edition and concept:
        num=re.search(r'\d+',concept);num=num[0] if num else ''
        # Limit concept matching to the SOURCE HEADING, not background keywords.
        candidates=[n for n in taxonomy['nodes'] if n['edition_key']==edition and n['node_type'] in ('chapter','unit') and
                    re.search(r'(?:Ind AS|Indian Accounting Standard)\s*'+re.escape(num)+r'\b',n['title'],re.I)]
        if len(candidates)==1:
            n=candidates[0];index={n['node_key']:n for n in taxonomy['nodes']};path=[]
            while n:path.append(n);n=index.get(n['parent_key'])
            official=[dict(key=n['node_key'],title=n['title'],type=n['node_type']) for n in reversed(path) if n['node_type']!='module']
            # A detailed exact phrase in the operative requirement can be proposed;
            # no incidental phrase from general scenario facts is a classification.
            text=question['content']['plain_text'];requirement=re.split(r'\bRequired\s*:|\bYou are required\b',text,flags=re.I)[-1] if re.search(r'\bRequired\s*:|\bYou are required\b',text,re.I) else '\n'.join(text.splitlines()[-8:])
            for n in taxonomy['nodes']:
                if n['edition_key']==edition and n['node_type'] in ('topic','subtopic') and len(n['title'])>12 and n['title'].lower() in requirement.lower() and n['node_key'].startswith(official[-1]['key']+'-'):
                    detail.append(dict(key=n['node_key'],title=n['title'],type=n['node_type']))
    return dict(historical_edition=edition,concept=concept,official_nodes=official+detail,
                method='manual override' if ledger else 'heading-supported' if heading and official else 'needs semantic owner review',
                reasoning='Printed preceding heading and edition-specific official contents were compared; detailed tags require owner confirmation.' if official else 'Historical chapter/topic path or question-level semantic analysis remains unresolved. No latest-edition substitute was made.',
                owner_correction=ledger,exact_topic_established=bool(detail))


def batches(rows,out,prefix,max_rows=25,max_bytes=1500000,key='records',extra=None):
    paths=[];current=[]
    def flush():
        if not current:return
        name=prefix+'-'+str(len(paths)+1).zfill(3)+'.json';payload={'batch_key':'FR-V05-'+prefix+'-'+str(len(paths)+1).zfill(3),key:list(current)}
        if extra:payload.update(extra)
        write(out/name,payload);paths.append(name);current.clear()
    for row in rows:
        if len(json.dumps(row,ensure_ascii=False).encode())>max_bytes:raise ValueError('One record exceeds safe import size; keep it in private detail and resolve before import')
        if current and (len(current)>=max_rows or len(json.dumps(current+[row],ensure_ascii=False).encode())>max_bytes):flush()
        current.append(row)
    flush();return paths


def workbook(rows,documents,taxonomy,exceptions,out,index,source_inventory=None):
    book=Workbook();book.remove(book.active)
    headers=['Stable ID','Source Type','Attempt','Series','Question Number','Subpart','Full Question Text','Full Answer Text','Chapter','Unit','Topic','Subtopic','Additional Tags','Source URL','Question Page','Answer Page','Question Type','Extraction Warnings','Classification Warnings','Review Decision','Owner Corrections','Notes','Revision','Source Hash','Question HTML SHA-256','Answer HTML SHA-256']
    sheets={name:book.create_sheet(name) for name in ['Source Inventory','Questions','Answers','Topic Hierarchy','Exceptions','Progress Summary','Import Manifest','Approval Decisions']}
    sheets['Questions'].append(headers);sheets['Answers'].append(['Question ID','Exact Question Reference','Subpart','Answer Source Hash','Answer Page','Full Answer Text','Mapping State'])
    sheets['Source Inventory'].append(['Drive ID','Filename','Role','Attempt','Issue Period','Series','PDF SHA-256','Pages','Outcome','Warnings','Official URL'])
    for d in source_inventory or []:sheets['Source Inventory'].append([d.get('file_id','not in Drive catalogue'),d.get('filename',''),d.get('source_kind','identity unconfirmed'),d.get('attempt',d.get('edition','')),d.get('issue_period',''),d.get('series',''),d.get('sha256','not retrieved'),d.get('pages',''),d.get('outcome',''),'; '.join(d.get('warnings',[])) or d.get('required_action',''),d.get('official_url',d.get('source_url',''))])
    sheets['Topic Hierarchy'].append(['Edition','Key','Parent','Type','Exact heading','Heading Number','Printed Page','Official/Editorial','Official URL','PDF Evidence','Review State'])
    for n in taxonomy['nodes']:sheets['Topic Hierarchy'].append([n['edition_key'],n['node_key'],n['parent_key'],n['node_type'],n['title'],n.get('heading_number',''),n.get('printed_page',''),n['heading_origin'],n['source_url'],n['source_location'],'Needs Review'])
    sheets['Exceptions'].append(['Stable ID / Source','Exception','Action']);sheets['Progress Summary'].append(['Measure','Value']);sheets['Progress Summary'].append(['Candidate question/subpart rows in this book',len(rows)]);sheets['Progress Summary'].append(['Academically verified pairs',0]);sheets['Progress Summary'].append(['Automatic publication',0])
    sheets['Import Manifest'].append(['Key','Value']);sheets['Import Manifest'].append(['format','CAFCC-FR-PRIVATE-1']);sheets['Import Manifest'].append(['instructions','Private authorised research. Use Questions columns Review Decision / Owner Corrections / Notes, or Approval Decisions; conflicting edits are rejected. All content remains Needs Review. Import the matching JSON first. Decisions are applied only by an authenticated administrator after WordPress validates exact IDs, revision 1, source and HTML hashes. A changed WordPress record needs a freshly exported workbook.']);
    decision_headers=['Stable ID','Revision','Source Hash','Owner Decision','Corrected Node IDs','Reviewer Notes','Question HTML SHA-256','Answer HTML SHA-256'];sheets['Approval Decisions'].append(decision_headers)
    for row in rows:
        r=row['import_record'];qtext=row['question']['plain_text'];atext=row['answer']['plain_text'] if row['answer'] else ''
        def fit(text,kind):return text if len(text)<=32767 else 'Full '+kind+' exceeds Excel cell limit; private details/'+r['stable_id']+'.html contains the complete text. No text was truncated.'
        labels={n['node_type']:n['title'] for n in r['ai_proposal'].get('primary_path',r['ai_proposal']['official_nodes'])};qh=sha(r['rich_content']['question']['body_html']);ah=sha(r['rich_content'].get('answer',{}).get('body_html',''))
        sheets['Questions'].append([r['stable_id'],r['source_type'],r['exam_attempt'],r['mtp_series'],r['reference'],r['subreference'],fit(qtext,'question'),fit(atext,'answer'),labels.get('chapter',''),labels.get('unit',''),labels.get('topic',''),labels.get('subtopic',''),'; '.join(r['ai_proposal'].get('additional_node_keys',[])),r['official_url'],r['question_page'],r['answer'].get('page',''),r['question_format'],'; '.join(r['blocking_exceptions']),r['reasoning'],'Needs Review','','',1,r['source_hash'],qh,ah])
        sheets['Answers'].append([r['stable_id'],r['reference'],r['subreference'],r['answer'].get('source_hash',''),r['answer'].get('page',''),fit(atext,'answer'),'Needs Review'])
        sheets['Approval Decisions'].append([r['stable_id'],1,r['source_hash'],'','','',qh,ah])
        for e in r['blocking_exceptions']:sheets['Exceptions'].append([r['stable_id'],e,'Inspect the private source and correct the draft in WordPress.'])
    for e in exceptions:sheets['Exceptions'].append([e.get('file_id',e.get('source_url','taxonomy')),e.get('code',''),e.get('required_action','Owner review / source delivery required')])
    for name,sheet in sheets.items():
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='102542');cell.font=Font(color='FFFFFF',bold=True)
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                # Never create a formula from source prose or reviewer text.
                if cell.data_type=='f':cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top');cell.protection=Protection(locked=not((name=='Approval Decisions' and cell.column in (4,5,6)) or (name=='Questions' and cell.column in (20,21,22))))
                if name=='Questions':cell.fill=PatternFill('solid',fgColor='E4F6F3')
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=62 if name=='Questions' and col[0].column in (7,8,18,19) else 28
        # Protection only guards accidental edits; the server enforces integrity.
        sheet.protection.sheet=True;sheet.protection.autoFilter=False
    for name,column in [('Questions','T'),('Approval Decisions','D')]:
        validation=DataValidation(type='list',formula1='"Needs Review,Publish,Reject"',allow_blank=True);validation.error='Choose Needs Review, Publish or Reject';validation.showErrorMessage=True;sheets[name].add_data_validation(validation);validation.add(column+'2:'+column+str(len(rows)+1))
    filename='FR-private-owner-review-'+str(index).zfill(3)+'.xlsx';book.save(out/filename);return filename


def build(root,output,previous_path=None):
    root=Path(root);out=Path(output).resolve()
    if (ROOT/'work').resolve() not in out.parents:raise ValueError('Private output must stay under ignored work/')
    if out.exists():raise ValueError('Choose a new release directory; never overwrite an owner review workbook')
    out.mkdir(parents=True,mode=0o700);docs=json.loads((root/'extracted/documents.json').read_text());taxonomy=json.loads((root/'fr-official-hierarchy.json').read_text())
    previous={('R-'+r['question_identity']):r for r in json.loads(Path(previous_path).read_text())} if previous_path else {}
    known={r['source_sha256']:r['official_url'] for r in previous.values()};known['1ff2d19f3d6ddb69d1bfb84158585f30a0f6d4ef4373c85cc697c793bc2a6408']='https://resource.cdn.icai.org/84615bos68187-p1.pdf'
    rows=[];mcq=[];mixed=[];exceptions=list(taxonomy['exceptions']);duplicates=[];seen=set();inventory=[]
    for d in docs:
        source=d['source'];ident=d['identity'];source_url=known.get(source['source_sha256'],'');inv=dict(file_id=source['file_id'],filename=source['filename'],source_kind=source['source_kind'],sha256=source['source_sha256'],pages=d['pages'],attempt=ident['attempt'],issue_period=ident['issue_period'],series=ident['series'],official_url=source_url,outcome=d['processing_outcome'],owner_review='needs_review',warnings=list(ident['warnings']))
        if not source_url:inv['warnings'].append('exact_official_pdf_url_not_yet_reconciled');exceptions.append(dict(file_id=source['file_id'],filename=source['filename'],code='exact_official_pdf_url_not_yet_reconciled',required_action='Confirm the exact ICAI PDF for these bytes; never substitute a private Drive hyperlink.'))
        inventory.append(inv)
        exceptions += [dict(file_id=source['file_id'],filename=source['filename'],**issue) for issue in d.get('exceptions',[])]
        for q in d['records']:
            if q['side']!='question':continue
            if q['stable_id'] in seen:duplicates.append(dict(stable_id=q['stable_id'],file_id=source['file_id'],code='duplicate_candidate_identity_held'));continue
            seen.add(q['stable_id'])
            if q['format']=='MCQ':mcq.append(dict(source=inv,record=q,previous_classification=previous.get(q['stable_id'],{}).get('classification'),owner_correction=previous.get(q['stable_id'],{}).get('ledger_event')));continue
            if q['format']!='SUBJECTIVE':mixed.append(dict(source=inv,record=q));continue
            candidates=match_answers(d,q,docs);cl=classification(d,q,taxonomy,previous);warnings=list(ident['warnings'])+list(q['warnings'])+list(q['content']['warnings'])
            if len(candidates)!=1:warnings.append('answer_absent_or_multiple_candidates')
            answer=candidates[0] if len(candidates)==1 else None
            if source['source_sha256']=='8bdb55ac0ffd3f01a96c5caac07326d4529d30acf8963094828a348e9b22c737' and q['reference']=='19':answer=None;warnings.append('FR_Jan2026_Answer19_corrigendum_unresolved_do_not_publish')
            if not cl['historical_edition']:warnings.append('historical_edition_not_established_no_current_edition_substitution')
            if not cl['exact_topic_established']:warnings.append('detailed_official_topic_or_subtopic_requires_academic_mapping')
            if not source_url:warnings.append('exact_official_pdf_url_not_yet_reconciled')
            warnings+=['owner_boundary_transcription_and_answer_identity_review_required','written_source_wide_permission_scope_not_entered']
            if answer:warnings+=answer['record']['content']['warnings']
            old=previous.get(q['stable_id']);source_key=old['document_key'] if old else 'FR-'+source['source_sha256'][:32]
            page=lambda r:str(r['start']['page'])+('–'+str(r['end_exclusive']['page']) if r['start']['page']!=r['end_exclusive']['page'] else '')
            answer_data={};content={'question':dict(body_html=q['content']['body_html'],author_name='The Institute of Chartered Accountants of India (ICAI)',source_location='PDF p.'+page(q),source_text_sha256=q['content']['source_text_sha256'],structure={},checks=q['content']['fidelity'])}
            if answer:
                ad=answer['document'];a=answer['record'];answer_data=dict(source_hash=ad['source']['source_sha256'],official_url=known.get(ad['source']['source_sha256'],''),reference=a['reference'],subreference=a['subreference'],page=page(a),question_identity=q['stable_id'],question_reference=q['reference'],question_subreference=q['subreference'],question_source_hash=source['source_sha256'],subject_code='FR',source_type=ident['source_type'],exam_attempt=ident['attempt'],mtp_series=ident['series'],issue_period=ad['identity']['issue_period'],notes='Provisional complete-document/section/reference candidate, not an academically verified answer match.',mode='licensed')
                content['answer']=dict(body_html=a['content']['body_html'],author_name='The Institute of Chartered Accountants of India (ICAI)',source_location='PDF p.'+page(a),source_text_sha256=a['content']['source_text_sha256'],structure={},checks=a['content']['fidelity'])
            attempt=ident['attempt'] or ('Unconfirmed target; issued '+ident['issue_period'] if ident['issue_period'] else 'Unconfirmed — inspect visible source')
            record=dict(stable_id=q['stable_id'],source_hash=source['source_sha256'],source_key=source_key,subject_code='FR',source_type=ident['source_type'],exam_attempt=attempt,issue_period=ident['issue_period'],mtp_series=ident['series'] or ('Unconfirmed — owner review' if ident['source_type']=='mtp' else ''),section=q['section'],reference=q['reference'],subreference=q['subreference'],title='ICAI Financial Reporting '+source['source_kind']+' · '+attempt,official_url=source_url,scheme='Historical scheme needs confirmation' if not cl['historical_edition'] else 'New Scheme — historical edition proposal pending owner review',edition_id=None,question_page=page(q),question_coordinates={'start':q['start'],'end_exclusive':q['end_exclusive'],'parent_shared_context_retained':q['parent_shared_context']},heading_evidence=q['heading'] or {},classification=[],reasoning=cl['reasoning'],confidence='Provisional; no academic accuracy percentage asserted',exceptions=sorted(set(warnings)),blocking_exceptions=sorted(set(warnings)),answer=answer_data,review_notes='',rights_mode='licensed_private_review',rights_basis='Owner expressly authorised private FR extraction and storage; owner reports written ICAI publication permission with lifetime duration. Publication scope/reference must be entered and confirmed privately before public use.',rights_grant_id=None,rich_content=content,extraction_status='needs_review',source_revision='SHA-256 '+source['source_sha256'],engine_revision={'version':d['engine'],'code_sha256':d['engine_hash']},previous_corrections=[cl['owner_correction']] if cl['owner_correction'] else [],ai_proposal={'historical_edition':cl['historical_edition'],'concept':cl['concept'],'official_nodes':[{'key':n['key'],'title':n['title'],'node_type':n['type']} for n in cl['official_nodes']],'method':cl['method']},marks=None,marks_verified=False,question_type='unknown',question_format='subjective')
            if answer_data:record['answer']['exam_attempt']=attempt;record['answer']['mtp_series']=record['mtp_series']
            row=dict(import_record=record,question=q['content'],answer=answer['record']['content'] if answer else None,answer_candidates=[dict(source_hash=a['document']['source']['source_sha256'],reference=a['record']['reference'],subreference=a['record']['subreference'],start=a['record']['start'],end_exclusive=a['record']['end_exclusive']) for a in candidates],source=inv)
            rows.append(row)
            detail='<html><meta charset="utf-8"><meta name="referrer" content="no-referrer"><title>Private '+html.escape(q['stable_id'])+'</title><style>body{font:17px system-ui;max-width:1100px;margin:2rem auto;padding:1rem}table{border-collapse:collapse}td,th{border:1px solid #889;padding:.4rem}section{margin:2rem 0}p{white-space:pre-wrap}</style><h1>PRIVATE · NEEDS REVIEW</h1><p>'+html.escape(source['filename']+' · '+attempt+' · Q'+q['reference']+' '+q['subreference'])+'</p><section><h2>Question</h2>'+q['content']['body_html']+'</section><section><h2>Provisional answer candidate</h2>'+(row['answer']['body_html'] if row['answer'] else '<p>No unique answer match established.</p>')+'</section><pre>'+html.escape(json.dumps({'classification':cl,'warnings':warnings,'question_boundary':record['question_coordinates'],'answer_candidates':row['answer_candidates']},ensure_ascii=False,indent=2))+'</pre></html>'
            (out/'details').mkdir(exist_ok=True);(out/'details'/(q['stable_id']+'.html')).write_text(detail)
    # Retain exact Drive metadata privately; never import it into a public resource.
    known_ids={d['source']['file_id'] for d in docs};extra=[]
    for s in json.loads((root/'drive-pdfs.json').read_text()):
        if s['id'] not in known_ids:extra.append(dict(file_id=s['id'],filename=s['title'],folder=s.get('folder_path',''),outcome='identity_and_third_party_rights_unconfirmed',required_action='Confirm publisher and separate rights. ICAI permission does not automatically cover teacher-authored books or notes.'))
    resource_manifest=json.loads((root/'study-pdf-inventory.json').read_text());study_inventory=resource_manifest['inventory'] if isinstance(resource_manifest,dict) else resource_manifest
    official=[]
    for entry in study_inventory:
        official.append(dict(source_url=entry['url'],filename=entry['url'].rsplit('/',1)[-1] if entry['url'] else entry['title'],source_kind='study_material',edition='May '+entry['year'],title=entry['title'],outcome='contents_metadata_recovered_only' if entry['title']=='Initial Pages' else 'raw_pdf_required_for_layout_extraction',required_action='Supply the authorised original PDF privately. Direct ICAI CDN download probe returned HTTP 403; these study URLs were individually inventoried through the web reader, not raw-downloaded.'))
    discovery=json.loads((root/'official-resource-discovery.json').read_text());raw='\n'.join(c['text'] for c in discovery['additional_pdfs']['content'] if c['type']=='text')
    for section in re.split('-'*80,raw):
        url=re.search(r'\((https://resource\.cdn\.icai\.org/[^)]+)\)',section)
        if url and url[1] not in known.values():official.append(dict(source_url=url[1],filename=url[1].rsplit('/',1)[-1],outcome='raw_pdf_required_for_layout_extraction',required_action='Save the official PDF to the private FR Drive folder; the web reader cannot supply raw coordinates/bytes for this source.'))
    for item in extra+official:exceptions.append(dict(file_id=item.get('file_id','not_in_drive_catalogue'),filename=item.get('filename',''),source_url=item.get('source_url',''),code=item['outcome'],required_action=item['required_action']))
    for row in rows:
        for warning in row['import_record']['blocking_exceptions']:exceptions.append(dict(stable_id=row['import_record']['stable_id'],file_id=row['source']['file_id'],filename=row['source']['filename'],code=warning,required_action='Private source/answer/classification review required; no publication approval is conveyed.'))
    write(out/'private-source-register.json',[dict(file_id=d['source']['file_id'],filename=d['source']['filename'],private_drive_url=d['source']['private_url'],source_sha256=d['source']['source_sha256'],title_evidence=d['identity']['title_evidence'],identification_warnings=d['identity']['warnings'],private_owner_review_only=True) for d in docs]);write(out/'source-inventory.json',inventory+extra+official);write(out/'fr-official-hierarchy.json',taxonomy);write(out/'subjective-candidates.json',rows);write(out/'preserved-mcq-research.json',mcq);write(out/'mixed-review-hold.json',mixed);write(out/'duplicate-exceptions.json',duplicates);write(out/'exceptions.json',exceptions)
    import_files=batches([r['import_record'] for r in rows],out/'imports','studio-drafts');hierarchy_files=batches(taxonomy['nodes'],out/'taxonomy-imports','official-hierarchy',max_rows=100,key='nodes',extra={'editions':taxonomy['editions']})
    books=[]
    for i in range(0,len(rows),150):books.append(workbook(rows[i:i+150],docs,taxonomy,[e for e in exceptions if not e.get('stable_id')],out,len(books)+1,inventory+extra+official))
    counts=Counter(r['import_record']['source_type'] for r in rows);summary=dict(timestamp=datetime.now(timezone.utc).isoformat(),catalogue_pdfs=79,pdfs_locally_processed=len(docs),pages=sum(d['pages'] for d in docs),source_failures=sum(d['processing_outcome']=='failed' for d in docs),rtp_subjective_candidates=counts['rtp'],mtp_subjective_candidates=counts['mtp'],exam_subjective_candidates_from_suggested_answer_books=counts['exam'],subjective_parent_records=sum(not r['import_record']['subreference'] for r in rows),subjective_subpart_records=sum(bool(r['import_record']['subreference']) for r in rows),mcq_candidates_separated=len(mcq),mixed_candidates_held=len(mixed),unique_answer_candidates=sum(r['answer'] is not None for r in rows),unresolved_answer_candidates=sum(r['answer'] is None for r in rows),academically_verified_question_answer_pairs=0,official_heading_nodes=len(taxonomy['nodes']),topic_proposals_with_heading_evidence=sum(bool(r['import_record']['ai_proposal']['official_nodes']) for r in rows),formatting_exception_records=sum(bool(r['question']['warnings'] or r['answer'] and r['answer']['warnings']) for r in rows),draft_wordpress_records=len(rows),owner_approvals=0,automatic_publications=0,additional_drive_metadata_only=len(extra),duplicate_candidate_identities_held=len(duplicates),official_additional_sources_inventoried=len(official),archive_acceptance='INCOMPLETE: external official sources, transcription, answer identity, historical classification and owner review remain outstanding')
    write(out/'coverage-report.json',summary);write(out/'import-order.json',{'first':'taxonomy-imports in numeric order (all new headings Draft; existing IDs/wording retained)','second':'imports in numeric order, Preview → inspect conflicts → Apply drafts','existing_record_conflicts':'Use the explicit enrichment workflow; never delete and reimport an approved candidate.','hierarchy_batches':hierarchy_files,'question_batches':import_files,'review_workbooks':books,'publication':'Owner review only; no record approved by this export.'})
    (out/'READ-FIRST.txt').write_text('PRIVATE FINANCIAL REPORTING REVIEW — VERSION 0.5\n\nThis is a substantive extraction candidate package, not a verified or complete question bank. Read coverage-report.json and exceptions.json. Every question and proposed answer is Needs Review. Raw source PDFs remain in the authorised Drive archive; no sharing settings were changed. Do not put this package in GitHub or public WordPress uploads.\n\nInstall the lightweight plugin ZIP separately. In Question Review Studio, import taxonomy batches first using the v5 hierarchy uploader, then research JSON batches. Existing candidate conflicts retain all owner corrections and IDs. Do not delete them; explicitly enrich missing private content or compare in Studio.\n\nThe full private review books use Approval Decisions for WordPress upload. Import the matching JSON first. Only an authenticated administrator can Validate and Apply; exact revision/source/content checks prevent stale decisions. Full formatted content is also in details/<Stable ID>.html. Excel cells longer than its limit reference those complete private details.\n\nThe written ICAI permission was not supplied as a file here. The owner stated lifetime duration. Enter and confirm the actual authority, reference/date, publisher, covered periods/material/use and attribution in the private rights register before publication. No reference/date, question approval or academic accuracy was invented.\n')
    write(out/'prior-owner-corrections.json',[r.get('ledger_event') for r in previous.values() if r.get('ledger_event')]);write(out/'before-after-comparison.json',[{'stable_id':r['import_record']['stable_id'],'previous':previous.get(r['import_record']['stable_id'],{}).get('classification'), 'current':r['import_record']['ai_proposal'],'boundary':r['import_record']['question_coordinates'],'review_status':'needs_review'} for r in rows if r['import_record']['stable_id'] in previous])
    checksums={str(p.relative_to(out)):sha(p.read_bytes()) for p in sorted(out.rglob('*')) if p.is_file()};write(out/'manifest.json',{'format':'CAFCC-FR-PRIVATE-1','review_status':'needs_review','publication_status':'draft','files':checksums,'counts':summary})
    return summary


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('root');p.add_argument('--output',required=True);p.add_argument('--previous');a=p.parse_args();print(json.dumps(build(a.root,a.output,a.previous),indent=2))
