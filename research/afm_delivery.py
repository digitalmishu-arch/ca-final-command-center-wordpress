"""One directly uploadable AFM collection, with private review evidence.

Requires the bounded collection adapter; every embedded child remains the
established Academic Package v1 format. Never approves or publishes records.
"""
from collections import Counter,defaultdict
from copy import deepcopy
from pathlib import Path
import argparse,html,json,re,zipfile,shutil,posixpath
from datetime import datetime,timezone
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Protection
from openpyxl.worksheet.datavalidation import DataValidation
import afm_reconcile as afm
import archive_catalog as cat
import gold_layout as gold

ROOT=Path(__file__).resolve().parents[1]
MCQ_KEYS='stable_id source_occurrence_id subject_code source_type exam_attempt mtp_series issue_period edition_key scheme section reference source_key source_hash official_url question_page question_location question_html question_checks question_structure options scenario answer correct_options key_policy key_verified explanation_html tags exceptions review_status academic_approval review_notes'.split()

def raw(obj):return (json.dumps(obj,ensure_ascii=False,separators=(',',':'))+'\n').encode()
def plain(s):return html.unescape(re.sub('<[^>]+>',' ',s or''))
def safe(s):return afm.slug(str(s))or'not-established'
def bounded_taxonomy(bank):
    """Fit the established 120-character key schema without changing titles."""
    mapping={};editions={key:key[:63]+'-'+gold.sha(key)[:16]for key in bank['taxonomies']if len(key)>80}
    for tax in bank['taxonomies'].values():
        for node in tax['nodes']:
            key=node['node_key']
            if len(key)>120:mapping[key]=key[:103]+'-'+gold.sha(key)[:16]
    if len(set(mapping.values()))!=len(mapping):raise ValueError('Taxonomy key collision')
    for tax in bank['taxonomies'].values():
        for edition in tax['editions']:edition['edition_key']=editions.get(edition['edition_key'],edition['edition_key'])
        for node in tax['nodes']:
            node['node_key']=mapping.get(node['node_key'],node['node_key']);node['parent_key']=mapping.get(node.get('parent_key',''),node.get('parent_key',''))
            if 'edition_key'in node:node['edition_key']=editions.get(node['edition_key'],node['edition_key'])
    for row in bank['subjective']:
        proposal=row['import_record']['ai_proposal']
        if 'historical_edition'in proposal:proposal['historical_edition']=editions.get(proposal['historical_edition'],proposal['historical_edition'])
        for node in proposal.get('primary_path',[]):node['node_key']=mapping.get(node['node_key'],node['node_key'])
        proposal['additional_node_keys']=[mapping.get(k,k)for k in proposal.get('additional_node_keys',[])]
    for row in bank.get('mcqs',[]):
        r=row['import_record'];r['edition_key']=editions.get(r['edition_key'],r['edition_key'])
    for source in bank.get('sources',{}).values():source['edition_key']=editions.get(source['edition_key'],source['edition_key'])
    bank['taxonomies']={editions.get(key,key):value for key,value in bank['taxonomies'].items()}
    return {**mapping,**editions}
def child_zip(path,key,attempt,edition,tax,sources,scenarios,records,kind):
    members={};items=[]
    def add(name,k,data):
        v=raw(data);assert len(v)<=2000000,(name,len(v));members[name]=v;items.append(dict(path=name,kind=k,bytes=len(v),sha256=gold.sha(v)))
    # Dependency order remains the existing format, bounded per request.
    for n in range(0,len(tax['nodes']),100):add('items/taxonomy_'+str(n//100)+'.json','taxonomy',dict(batch_key=tax['batch_key']+'-'+str(n//100),editions=tax['editions'],nodes=tax['nodes'][n:n+100]))
    add('items/sources.json','sources',dict(sources=list(sources.values())))
    if scenarios:add('items/scenarios.json','scenarios',dict(batch_key=key+'-SC',scenarios=list(scenarios.values())))
    for n in range(0,len(records),10):
        batch=records[n:n+10];obj=dict(format='CAFCC-MCQ-DRAFTS-1',records=batch)if kind=='mcqs'else dict(batch_key=key+'-'+str(n//10),records=batch)
        add('items/records_'+str(n//10).zfill(3)+'.json',kind,obj)
    m=dict(format='CAFCC-ACADEMIC-PACKAGE-1',package_key=key,version='0.8.2',created_at=datetime.now(timezone.utc).isoformat(),scope=dict(subject_code='AFM',exam_attempt=attempt,edition_key=edition),coverage=dict(draft_candidates=len(records),academic_approvals=0,publications=0),items=items)
    members['manifest.json']=raw(m);assert len(members)<=65
    path.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED,compresslevel=9)as z:
        for name,content in members.items():z.writestr(name,content)
    assert path.stat().st_size<=8000000
    return m

def workbook(path,bank,details):
    wb=Workbook();wb.remove(wb.active);headers=['Question ID','Source Type','Year','Attempt / Source Period','MTP Series','Original Subject','Source PDF','Question Number','Full Question Text','Full Answer Text / Key','Chapter / Editorial Crosswalk','Topic','Subtopic / Additional Concepts','Original PDF','Answer PDF','Question Page','Answer Page','Match State','Exact Issue / Evidence','Status','Owner Correction / Decision','Notes','Detail File','Source SHA256','Revision']
    sheets={name:wb.create_sheet(name)for name in ['Master AFM Question Bank','PYP Questions','RTP Questions','MTP Questions','MCQs','Answer Matching Issues','Missing Answers','Question Boundary Issues','Tables Formula Issues','Topic Syllabus Mapping Issues','Source PDF Identity Issues','Duplicate Repeated Questions','Source Coverage Summary','Pattern Groups']}
    for name,sheet in sheets.items():
        if name not in('Source Coverage Summary','Pattern Groups'):sheet.append(headers)
    source_index={r['file_id']:r for r in bank['source_inventory']};patterns=defaultdict(list)
    for kind,rows in [('subjective',bank['subjective']),('mcq',bank['mcqs'])]:
        for row in rows:
            r=row['import_record'];src=source_index[row['source_file_id']];sid=r['stable_id'];question=row['question']['plain_text'];answer=(row.get('answer')or{}).get('plain_text','')if kind=='subjective'else'Official key: '+', '.join(r['correct_options'])if r.get('correct_options')else'ANSWER PENDING'
            if kind=='subjective'and not answer:
                answer='ANSWER PENDING'
                if row.get('answer_candidate'):answer+='\nUNASSIGNED PARENT ANSWER CANDIDATE — NOT A MATCHED CHILD SOLUTION:\n'+row['answer_candidate']['plain_text']
            if kind=='subjective':
                sk=r.get('practice_structure',{}).get('scenario_key');common=bank['scenarios'].get(sk,{}).get('body_html','');question=(plain(common)+'\n\n'if common else'')+question
            else:
                if r.get('scenario'):question=plain(r['scenario']['html'])+'\n\n'+question
                question+='\n\n'+'\n'.join(o['id']+'. '+plain(o['html'])for o in r['options'])
                if r.get('explanation_html'):answer+='\n\nSource explanation:\n'+plain(r['explanation_html'])
            issues=r.get('blocking_exceptions',r.get('exceptions',[]));proposed=row.get('proposed_topics',[]);topic=proposed[0]['label']if proposed else', '.join(t['label']for t in r.get('tags',[]))or'TOPIC REVIEW REQUIRED';chapter=proposed[0]['chapter_title']if proposed else'';year=(re.search(r'20\d{2}',r['exam_attempt'])or re.search(r'20\d{2}',src['filename']));year=year[0]if year else'Unestablished'
            cells=[sid,r['source_type'],year,r['exam_attempt'],r.get('mtp_series',''),row['original_subject'],src['filename'],r['reference']+('('+r['subreference']+')'if r.get('subreference')else''),question,answer,chapter,topic,', '.join(p['label']for p in proposed[1:]),r['official_url'],r.get('answer',{}).get('official_url',''),r.get('question_page',''),r.get('answer',{}).get('page',''),row.get('match_state','Official: No correct option — unscored'if r.get('answer',{}).get('key_evidence',{}).get('no_correct_option')else'Explicit source key'if r.get('key_verified')else'Key correspondence unresolved'),'\n'.join(issues),'Needs Review','','',details.get(sid,''),r['source_hash'],1]
            if not cells[14]and row.get('answer_file_id'):cells[14]='https://drive.google.com/file/d/'+row['answer_file_id']+'/view'
            if not cells[16]and row.get('answer_candidate_page'):cells[16]=row['answer_candidate_page']
            # Full source prose is retained in details/JSON, never silently cut.
            for c in (8,9):
                if len(str(cells[c]))>32767:cells[c]=str(cells[c])[:30000]+'\n[EXCEL LIMIT: full content is in '+details[sid]+']'
            destinations=['Master AFM Question Bank']+(['MCQs']if kind=='mcq'else({'exam':['PYP Questions'],'rtp':['RTP Questions'],'mtp':['MTP Questions']}.get(r['source_type'],[])))
            if kind=='subjective'and row['match_state']in('B','C','E'):destinations.append('Answer Matching Issues')
            if kind=='subjective'and row['match_state']=='D'or kind=='mcq'and not r.get('key_verified')and not r.get('answer',{}).get('key_evidence'):destinations.append('Missing Answers')
            if kind=='mcq'and not r.get('key_verified')and r.get('answer',{}).get('key_evidence')and not r['answer']['key_evidence'].get('no_correct_option'):destinations.append('Answer Matching Issues')
            if any(re.search('boundary|part_sequence|alternative_identity|number_sequence',x,re.I)for x in issues):destinations.append('Question Boundary Issues')
            if any(re.search('table|formula|glyph|numeric|geometry|font',x,re.I)for x in issues):destinations.append('Tables Formula Issues')
            if any(re.search('TOPIC_REVIEW|historical|syllabus',x,re.I)for x in issues):destinations.append('Topic Syllabus Mapping Issues')
            if any(re.search('source_identity|target_attempt',x,re.I)for x in issues):destinations.append('Source PDF Identity Issues')
            for name in set(destinations):sheets[name].append(cells)
            for issue in issues:
                group=re.sub(r'_page_\d+','_page_N',issue);patterns[group].append(sid)
    # Only byte/content-copy evidence is listed here. No unreviewed repeat claim.
    for copy in bank['source_copies']:
        src=source_index[copy['file_id']];canon=source_index[copy['canonical_file_id']];sheets['Duplicate Repeated Questions'].append(['SOURCE-COPY:'+copy['file_id'],src['identity']['source_type'],'',src['identity']['exam_attempt'],src['identity']['series'],src['identity']['actual_subject'],src['filename'],'','','','','','','https://drive.google.com/file/d/'+src['file_id']+'/view','',1,'','Source copy',copy['reason']+' Canonical: '+canon['file_id'],'Needs Review','','','',src['source_hash'],1])
    for issue in bank.get('source_issues',[]):
        src=source_index.get(issue.get('file_id'),{})
        if not src:continue
        d=src['identity'];sheets['Source PDF Identity Issues'].append(['SOURCE:'+src['file_id'],d['source_type'],'',d['exam_attempt'],d.get('series',''),d.get('actual_subject',''),src['filename'],'','','','','','','https://drive.google.com/file/d/'+src['file_id']+'/view','',issue.get('page',''),'','Source-level issue',issue['reason'],'Needs Review','','','',src['source_hash'],1])
    coverage=bank['coverage'];ch=['Drive File ID','Filename','Pages','SHA256','Actual Subject','Scheme','Source Type','Attempt / Source Period','Issue Period','Series','Role','Subjective Parents','Subjective Practice Parts','MCQ Candidates','Source-corroborated Answer Matches','Answer Candidates','Answer Pending','Blocked Question Drafts','Subjective Import Drafts','MCQ Import Drafts','Outcome / Exclusion','PDF URL','Owner Correction','Solutions supplied to question sources','Unresolved supplied answer contexts']
    s=sheets['Source Coverage Summary'];s.append(ch)
    for r in coverage:s.append([r['file_id'],r['filename'],r['pages'],r['sha256'],r['historical_subject'],r['scheme'],r['source_type'],r['attempt'],r['issue_period'],r['series'],r['role'],r['subjective_parents'],r['subjective_parts'],r['mcq_candidates'],r['source_corroborated_answers'],r['answer_candidates'],r['answer_pending'],r['blocked_questions'],r['import_subjective'],r['import_mcq'],r['status']+'; '+'; '.join(r['exclusion_reasons']),r['source_url'],'',r.get('provided_answer_associations',0),r.get('provided_unresolved_contexts',0)])
    s=sheets['Pattern Groups'];s.append(['Actual Issue Pattern','Affected Records','Question IDs (see detail tabs for evidence)','Owner Rule / Decision'])
    for key,ids in sorted(patterns.items(),key=lambda x:-len(x[1])):s.append([key,len(set(ids)),', '.join(sorted(set(ids)))[:32000],''])
    if bank.get('r011_warning_ledger'):
        s=wb.create_sheet('r010 Warning Reconciliation');s.append(['Old Question ID','r011 Question ID','Old Warning State','r011 Source Match State','Audit Outcome','Question PDF','Answer PDF','Question Page','Answer Page','Exact Evidence','Owner Correction / Decision'])
        for row in bank['r011_warning_ledger']:
            qr=row.get('question_region')or{};ar=row.get('answer_region')or{};s.append([row['old_id'],row['new_id'],row['old_state'],row['new_state'],row['outcome'],'https://drive.google.com/file/d/'+row['source_file_id']+'/view','https://drive.google.com/file/d/'+row['answer_file_id']+'/view',qr.get('start',{}).get('page',''),ar.get('start',{}).get('page',''),json.dumps(row['evidence'],ensure_ascii=False),''])
        s=wb.create_sheet('Structure ID Reconciliation');s.append(['Old Question ID','r011 Question ID','Action','Original Part','r011 Part','Student Activity Policy','Existing Edit Policy','Owner Decision'])
        for row in bank['r011_reconciliation']:s.append([row['old_id'],row['new_id'],row['action'],row['old_subreference'],row['new_subreference'],row['student_history'],row['existing_edit_policy'],''])
    for sheet in wb:
        sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.zoomScale=85
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173E53');cell.alignment=Alignment(wrap_text=True,vertical='top')
        sheet.row_dimensions[1].height=35
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top');cell.protection=Protection(locked=cell.column not in(21,22))
                if isinstance(cell.value,str)and cell.value.startswith('https://drive.google.com/'):cell.hyperlink=cell.value;cell.font=Font(color='007B80',underline='single')
            sheet.row_dimensions[row[0].row].height=80
        for col in sheet.columns:
            idx=col[0].column;sheet.column_dimensions[col[0].column_letter].width=48 if idx in(9,10,19)else 28 if idx in(1,7,14,15,21,22,23,24)else 19
        if sheet.title not in('Source Coverage Summary','Pattern Groups'):
            dv=DataValidation(type='list',formula1='"Needs Review,Publish,Reject,Hold"');sheet.add_data_validation(dv);dv.add('U2:U'+str(max(2,sheet.max_row)))
    wb.save(path)

def build(bank_path,output,revision,evidence=None):
    bank=json.loads(Path(bank_path).read_text());key_mapping=bounded_taxonomy(bank);out=cat.private(output);out.mkdir(parents=True,exist_ok=True);base=out/'AFM';base.mkdir(exist_ok=True);details={};import_order=[]
    cat.write(base/'integrity/taxonomy-key-reconciliation.json',key_mapping)
    for row in bank['subjective']+bank['mcqs']:
        r=row['import_record'];src=next(s for s in bank['source_inventory']if s['file_id']==row['source_file_id']);legacy='Historical-SFM'if row['original_subject']=='SFM'else'Current-AFM';kind={'exam':'PYP','rtp':'RTP','mtp':'MTP','other':'CASE-BOOKLETS'}[r['source_type']];year=(re.search(r'20\d{2}',r['exam_attempt'])or re.search(r'20\d{2}',src['filename']));year=year[0]if year else'unestablished';folder=base/('mcq'if'options'in r else'question-answer-html')/legacy/kind/year/safe(r['exam_attempt'])/safe(src['filename'])
        folder.mkdir(parents=True,exist_ok=True);path=folder/(r['stable_id']+'.html');details[r['stable_id']]=path.relative_to(base).as_posix()
        common=bank['scenarios'].get(r.get('practice_structure',{}).get('scenario_key'),{}).get('body_html','')if'options'not in r else(r.get('scenario')or{}).get('html','');question=r.get('question_html',r.get('rich_content',{}).get('question',{}).get('body_html',''));answer=r.get('rich_content',{}).get('answer',{}).get('body_html','')or'<p>ANSWER PENDING</p>'
        if row.get('answer_candidate'):answer+='<h3>Unassigned parent answer candidate — NOT the matched child solution</h3><p>Candidate PDF page '+str(row.get('answer_candidate_page',''))+'</p>'+row['answer_candidate']['body_html']
        for preview in row.get('formula_previews',[]):
            url=posixpath.relpath(preview['path'],path.relative_to(base).parent.as_posix())
            answer+='<p><a href="'+html.escape(url,quote=True)+'">Private original '+html.escape(preview['side'])+' formula crop · PDF '+str(preview['page'])+'</a> — source-layout correction required.</p>'
        if'options'in r:question+=''.join('<p><b>'+o['id']+'</b> '+o['html']+'</p>'for o in r['options']);answer='<p>Explicit source key: '+html.escape(', '.join(r['correct_options'])or'ANSWER PENDING')+'</p>'+r.get('explanation_html','')
        warnings=r.get('blocking_exceptions',r.get('exceptions',[]));path.write_text('<!doctype html><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Private AFM Needs Review</title><style>:root{color-scheme:light dark}body{font:17px/1.65 system-ui;max-width:1050px;margin:24px auto;padding:0 16px}table{border-collapse:collapse;max-width:100%;display:block;overflow-x:auto}td,th{border:1px solid #82909c;padding:8px}th{font-weight:700}a{color:#008f98}pre{white-space:pre-wrap}</style><h1>Needs Review — '+html.escape(r['stable_id'])+'</h1><p>'+html.escape(row['original_subject']+' · '+r['source_type']+' · '+r['exam_attempt']+' · '+r.get('mtp_series','')+' · Q'+r['reference']+('('+r['subreference']+')'if r.get('subreference')else''))+'</p><p><a href="'+r['official_url']+'">Original PDF</a> · PDF '+str(r['question_page'])+'</p><h2>Common information</h2>'+common+'<h2>Question</h2>'+question+'<h2>Answer / source-key candidate</h2>'+answer+'<p><a href="'+r.get('answer',{}).get('official_url','')+'">Answer PDF</a> · PDF '+str(r.get('answer',{}).get('page',''))+'</p><h2>Actual exceptions</h2><pre>'+html.escape('\n'.join(warnings))+'</pre>')
    for name,data in [('source-manifest/source-register.json',bank['source_inventory']),('coverage-reports/source-by-source.json',bank['coverage']),('coverage-reports/answer-matching.json',[dict(question_id=x['import_record']['stable_id'],state=x['match_state'],source_file_id=x['source_file_id'],answer_file_id=x['answer_file_id'],evidence=x['match_evidence'])for x in bank['subjective']]),('integrity/stable-id-reconciliation.json',bank['id_reconciliation']),('source-manifest/source-copies.json',bank['source_copies'])]:cat.write(base/name,data)
    wb=base/'academic-review/AFM-Owner-Review.xlsx';wb.parent.mkdir(exist_ok=True);workbook(wb,bank,details)
    groups=defaultdict(list)
    for kind,rows in [('questions',bank['subjective']),('mcqs',bank['mcqs'])]:
        for row in rows:
            if not row['import_eligible']:continue
            r=row['import_record'];groups[(row['source_file_id'],r['exam_attempt'],r['edition_key']if kind=='mcqs'else r['ai_proposal']['historical_edition'],kind)].append(r)
    for n,((fid,attempt,edition,kind),rs)in enumerate(groups.items(),1):
        for part in range(0,len(rs),60):
            selected=deepcopy(rs[part:part+60]);sourcekeys=set();scenario_keys=set()
            for r in selected:
                sourcekeys.add(r['source_key'])
                if r.get('answer',{}).get('source_key'):sourcekeys.add(r['answer']['source_key'])
                ah=r.get('answer',{}).get('source_hash')
                for sk,s in bank['sources'].items():
                    if s['source_hash']==ah:sourcekeys.add(sk)
                if kind=='mcqs':
                    for key in list(r):
                        if key not in MCQ_KEYS:del r[key]
                else:
                    # The established private provenance schema accepts IDs
                    # only. Historical identity remains in title/scheme and
                    # the private source register, never disguised as current.
                    r['source_provenance']={k:v for k,v in r.get('source_provenance',{}).items()if k in('drive_file_id','answer_drive_file_id')}
                    # Studio's existing answer-reference schema deliberately
                    # accepts strings (apart from coordinates/explicit IDs).
                    # Convert the PDF page scalar, preserving its exact value.
                    if r.get('answer')and 'page'in r['answer']:r['answer']['page']=str(r['answer']['page'])
                    sk=r.get('practice_structure',{}).get('scenario_key')
                    if sk:scenario_keys.add(sk)
            source_rows={k:v for k,v in bank['sources'].items()if k in sourcekeys};sc={k:v for k,v in bank['scenarios'].items()if k in scenario_keys};key='AFM-ALL-'+revision+'-'+str(n).zfill(3)+'-'+str(part//60+1)
            src=next(s for s in bank['source_inventory']if s['file_id']==fid);label=safe(src['filename'])[:60];filename=key+'-'+label+'-'+kind+'.zip';path=base/'import'/filename
            child_zip(path,key,attempt,edition,bank['taxonomies'][edition],source_rows,sc,selected,kind);import_order.append(dict(path=path.relative_to(out).as_posix(),bytes=path.stat().st_size,sha256=gold.sha(path.read_bytes()),records=len(selected),kind=kind,source_file_id=fid,source_filename=src['filename'],attempt=attempt))
    assert len(import_order)<=100,len(import_order)
    summary=dict(source_file_copies=len(bank['source_inventory']),unique_pdf_sha256=len({s['source_hash']for s in bank['source_inventory']}),native_content_copies=len(bank['source_copies']),subjective_candidates=len(bank['subjective']),mcq_candidates=len(bank['mcqs']),subjective_draft_imports=sum(x['records']for x in import_order if x['kind']=='questions'),mcq_draft_imports=sum(x['records']for x in import_order if x['kind']=='mcqs'),answer_matching=dict(Counter(x['match_state']for x in bank['subjective'])),child_packages=len(import_order),academic_approvals=0,publications=0)
    cat.write(base/'coverage-reports/summary.json',summary);cat.write(base/'import-order.json',import_order)
    (base/'READ-FIRST.txt').write_text('AFM/SFM PRIVATE REVIEW AND IMPORT — '+revision+'\n\nInstall the reviewed v0.8.2 application patch first. Upload this entire outer ZIP once through Upload Questions → Preview → Save as Drafts. It is a draft-only subject collection, not a WordPress plugin. It contains bounded existing Academic Package v1 children. Interrupted imports resume their checkpoints; reimport does not overwrite existing corrections. No question, MCQ or rights grant is automatically published.\n\nTo read the workbook on iPad: download this ZIP to Files, Uncompress, open AFM/academic-review/AFM-Owner-Review.xlsx in Excel or Numbers. Full selectable HTML is in the matching source/year/attempt folders. Issue sheets contain actual flags only. The large workbook is a research overview, not an automatic approval importer; use current WordPress decision exports after import for revision-safe decisions.\n\nHistorical SFM keeps its true paper identity. Current AFM crosswalks and detailed search concepts are editorial proposals. Unknown MTP targets are explicitly labelled as issue periods with target pending; they are not invented examination attempts. Source-corroborated matching is automated evidence, not independent academic verification. All real decisions are Needs Review.\n\nThe original January FR pilot and the other five subjects are not part of this import. Do not uninstall or reset the plugin.\n')
    if bank.get('r011_reconciliation'):
        with (base/'READ-FIRST.txt').open('a')as guide:
            guide.write('\nR011 STRUCTURAL CORRECTION\nSelect the exact r011 collection in Question Library. Existing r010 records, owner edits and student activity are retained, never reassigned. Changed existing IDs produce explicit import conflicts; this draft import does not apply those revisions over local edits. The private correction plan identifies each proposed replacement and its r010 HTML fingerprints. Reload the current question revision in Question Library/Studio, compare your edits, and save corrections using the existing revision-safe controls. Superseded nested IDs remain in the historical r010 dataset with an explicit old-to-new map; they are not silently deleted or used to transfer notes/timer history.\n\nOriginal numbered requirements are grouped into complete problems where the source supports it. Genuine separate Roman-numbered cases remain independent. All records stay Needs Review; unresolved material source/answer/table issues have blocking integrity flags. Known official No correct option MCQs are retained for unscored review, not treated as missing key extraction.\n')
    if evidence:
        eroot=cat.private(evidence)
        for p in eroot.rglob('*'):
            if not p.is_file():continue
            relative=p.relative_to(eroot)
            if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_/-]*\.(?:json|txt|csv|png|html)',relative.as_posix())or p.stat().st_size>8000000:raise ValueError('Unsafe private evidence attachment '+str(relative))
            destination=base/relative
            if destination.exists():
                if destination.read_bytes()==p.read_bytes():continue
                raise ValueError('Evidence cannot replace a generated delivery file')
            destination.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(p,destination)
    paths=[p for p in base.rglob('*')if p.is_file()];packages=[];attachments=[]
    for path in sorted(paths):
        v=dict(path=path.relative_to(out).as_posix(),bytes=path.stat().st_size,sha256=gold.sha(path.read_bytes()))
        (packages if path.parent==base/'import'else attachments).append(v)
    manifest=dict(format='CAFCC-ACADEMIC-COLLECTION-1',package_key='AFM-PAPER2-ALL-ATTEMPTS-'+revision,version='0.8.2',created_at=datetime.now(timezone.utc).isoformat(),subject_code='AFM',coverage=summary,packages=packages,attachments=attachments)
    archive=out/('AFM_PAPER_2_ALL_ATTEMPTS_REVIEW_AND_IMPORT_'+revision+'.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=9)as z:
        z.writestr('manifest.json',raw(manifest))
        for path in sorted(paths):z.write(path,path.relative_to(out).as_posix())
    assert archive.stat().st_size<=16000000,archive.stat().st_size
    print(json.dumps(dict(filename=archive.name,bytes=archive.stat().st_size,sha256=gold.sha(archive.read_bytes()),summary=summary),indent=2))
    return archive

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('bank');p.add_argument('--output',required=True);p.add_argument('--revision',default='r001');p.add_argument('--evidence');a=p.parse_args();build(a.bank,a.output,a.revision,a.evidence)
