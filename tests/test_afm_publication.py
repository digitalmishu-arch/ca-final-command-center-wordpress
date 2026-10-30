import sys
import unittest
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'research'))
from afm_publication import grid_valid, numeric_tokens
from afm_delivery import compact_xlsx
from openpyxl import Workbook, load_workbook


class PublicationCorrections(unittest.TestCase):
    def test_merged_grid_requires_full_coverage(self):
        model = dict(rows=3, columns=2, cells=[
            dict(row=0, column=0, rowspan=3, colspan=1, text='Authored scenario'),
            *[dict(row=r, column=1, rowspan=1, colspan=1, text=str(17+r)) for r in range(3)]])
        self.assertTrue(grid_valid(model))
        model['cells'][0]['rowspan'] = 2
        self.assertFalse(grid_valid(model))

    def test_overlap_and_overflow_are_rejected(self):
        self.assertFalse(grid_valid(dict(rows=1, columns=1, cells=[dict(row=0, column=0, rowspan=2)])))
        cell = dict(row=0, column=0, rowspan=1, colspan=1)
        self.assertFalse(grid_valid(dict(rows=1, columns=1, cells=[cell, cell])))

    def test_numerical_tokens_ignore_markup_but_keep_amounts(self):
        self.assertEqual(numeric_tokens('<p>17.50 and 2,600</p>'), numeric_tokens('<td>17.50</td><td>2,600</td>'))
        self.assertNotEqual(numeric_tokens('17.50'), numeric_tokens('17.05'))
        self.assertNotEqual(numeric_tokens('17 17'), numeric_tokens('17'))
        self.assertNotEqual(numeric_tokens('-17'), numeric_tokens('17'))
        self.assertNotEqual(numeric_tokens('(17.50)'), numeric_tokens('17.50'))
        self.assertNotEqual(numeric_tokens('17%'), numeric_tokens('17'))
        self.assertNotEqual(numeric_tokens('₹17'), numeric_tokens('$17'))

    def test_workbook_compaction_preserves_values_and_links(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'authored.xlsx'
            book = Workbook()
            for name in ['Authored', 'Repeated']:
                sheet = book.create_sheet(name)
                sheet.append(['Original authored fixture', '₹ 17.50', -31, None])
                sheet.append(['Original authored fixture', 'Private review only'])
                sheet['A1'].hyperlink = 'https://www.icai.org/'
                sheet.freeze_panes = 'B2'
            book.save(path)
            compact_xlsx(path)
            reopened = load_workbook(path)
            for name in ['Authored', 'Repeated']:
                self.assertEqual(reopened[name]['B1'].value, '₹ 17.50')
                self.assertEqual(reopened[name]['C1'].value, -31)
                self.assertEqual(reopened[name]['A1'].hyperlink.target, 'https://www.icai.org/')
                self.assertEqual(reopened[name].freeze_panes, 'B2')
