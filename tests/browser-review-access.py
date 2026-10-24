"""Actual local Studio/XLSX/access screens. Only original synthetic fixtures."""
from pathlib import Path
from urllib.parse import urlparse,parse_qs
import json,re,shutil,subprocess,uuid
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
BASE='http://127.0.0.1:8092';ROOT=Path(__file__).resolve().parents[1]
def login(page,user):
 page.goto(BASE+'/wp-login.php',wait_until='networkidle');page.locator('#user_login').fill(user);page.locator('#user_pass').fill('cafcc-disposable-only');page.locator('#wp-submit').click();page.wait_for_load_state('networkidle')
with sync_playwright() as p:
 browser=p.chromium.launch(headless=True,**({'executable_path':shutil.which('chromium')} if shutil.which('chromium') else {}),args=['--no-sandbox','--disable-dev-shm-usage'])
 admin=browser.new_page(viewport={'width':1024,'height':1366});errors=[];admin.on('pageerror',lambda e:errors.append(str(e)));login(admin,'test_admin');admin.goto(BASE+'/wp-admin/admin.php?page=cafcc-studio',wait_until='networkidle')
 expect(admin.locator('#studio-rows tr')).to_have_count(50);expect(admin.locator('#studio-pages')).to_contain_text('50 per page')
 admin.screenshot(path=str(ROOT/'docs/previews/v04-studio-ipad.png'),full_page=True)
 admin.locator('#studio-search').fill('TEST-ORIGINAL-OWNED');expect(admin.locator('#studio-rows tr')).to_have_count(1)
 admin.locator('#studio-rows tr td button').first.click();expect(admin.locator('#studio-panel')).to_be_visible();expect(admin.locator('#studio-detail')).to_contain_text('Authorised question preview');expect(admin.locator('#studio-detail')).to_contain_text('UNAPPROVED-NEW-SEARCH-WORD')
 admin.get_by_label('My private notes',exact=True).fill('Browser original private review note')
 admin.get_by_role('button',name='Save Corrections',exact=True).click();expect(admin.locator('#studio-detail')).to_contain_text('Authorised question preview')
 admin.locator('#studio-panel').evaluate('e=>e.scrollTop=0');admin.screenshot(path=str(ROOT/'docs/previews/v04-review-panel-ipad.png'),full_page=True)
 admin.get_by_role('button',name='Close',exact=True).click();admin.locator('#studio-rows input[type=checkbox]').check()
 try:
  with admin.expect_download(timeout=10000) as download:admin.get_by_role('button',name='Download Review Excel',exact=True).click()
 except Exception as error:raise RuntimeError('Download feedback: '+admin.locator('#studio-feedback').inner_text()) from error
 path=ROOT/'work/browser-private-review.xlsx';download.value.save_as(path);book=load_workbook(path);assert len(book.sheetnames)==8
 book['Approval Decisions']['F2']='Browser Excel correction note';edited=ROOT/'work/browser-private-decisions.xlsx';book.save(edited)
 admin.locator('#studio-excel').set_input_files(edited);expect(admin.locator('#studio-upload-preview')).to_be_visible();apply=admin.get_by_role('button',name='Apply validated decisions / draft batch',exact=True);expect(apply).to_be_enabled();apply.click();expect(admin.locator('#studio-upload-preview')).not_to_be_visible()
 admin.locator('#studio-excel').set_input_files([]);admin.locator('#studio-excel').set_input_files(edited);expect(admin.locator('#studio-feedback')).to_contain_text('current WordPress-issued')
 # Excel can be optional: direct Hold/review controls remain usable.
 admin.locator('#studio-rows tr td button').first.click();expect(admin.get_by_label('My private notes',exact=True)).to_have_value('Browser Excel correction note');admin.locator('[data-owner-action=needs_review]').click();expect(admin.locator('#studio-detail small')).to_contain_text('needs_review');admin.get_by_role('button',name='Close',exact=True).click()
 for width in [390,768,1024]:
  admin.set_viewport_size({'width':width,'height':900});assert admin.evaluate('document.documentElement.scrollWidth<=innerWidth'),width
 # Actual frontend registration -> email POST confirmation -> pending -> manual admin grant.
 student_context=browser.new_context(viewport={'width':390,'height':844});student=student_context.new_page();student.on('pageerror',lambda e:errors.append(str(e)))
 email='browser-'+uuid.uuid4().hex[:10]+'@example.invalid';student.goto(BASE+'/?cafcc_view=register',wait_until='networkidle');student.get_by_label('Your name',exact=True).fill('Browser New Student');student.get_by_label('Email address',exact=True).fill(email);student.get_by_label('Password',exact=True).fill('browser-private-only-password');student.get_by_role('button',name='Create account',exact=True).click();expect(student.locator('body')).to_contain_text('check your email')
 container=(ROOT/'work/preview-containers.txt').read_text().splitlines()[0]
 raw=subprocess.check_output(['docker','exec',container,'php','-r',"require '/var/www/html/wp-load.php';if(DB_NAME!=='cafcc_test')exit(1);echo json_encode(get_option('cafcc_test_last_mail'));"],text=True);mail=json.loads(raw);assert mail['to']==email or mail['to']==[email]
 url=re.search(r'https?://[^\s]+',mail['message']).group(0);assert parse_qs(urlparse(url).query)['cafcc_token']
 student.goto(url,wait_until='networkidle');student.get_by_role('button',name='Confirm Email',exact=True).click();expect(student.locator('body')).to_contain_text('awaiting administrator approval')
 student.goto(BASE+'/?cafcc_view=login',wait_until='networkidle');student.get_by_label('Email address',exact=True).fill(email);student.get_by_label('Password',exact=True).fill('browser-private-only-password');student.get_by_role('button',name='Sign in',exact=True).click();expect(student.locator('body')).to_contain_text('pending_approval');assert not student.locator('[data-cafcc-explorer]').count()
 student.screenshot(path=str(ROOT/'docs/previews/v04-pending-phone.png'),full_page=True)
 admin.goto(BASE+'/wp-admin/admin.php?page=cafcc-access',wait_until='networkidle');row=admin.locator('#access-rows tr').filter(has_text=email);expect(row).to_contain_text('Email verified');row.get_by_role('button',name='Apply manual decision',exact=True).click();expect(admin.locator('#access-feedback')).to_contain_text('recorded')
 admin.set_viewport_size({'width':1024,'height':1366});admin.screenshot(path=str(ROOT/'docs/previews/v04-access-ipad.png'),full_page=True)
 student.goto(BASE+'/?cafcc_view=explorer',wait_until='networkidle');expect(student.locator('[data-cafcc-explorer]')).to_be_visible()
 for view,text in [('resources','ICAI Resources'),('trends','Coming Soon'),('planner','Coming Soon'),('membership','Access is granted manually')]:
  student.goto(BASE+'/?cafcc_view='+view,wait_until='networkidle');expect(student.locator('body')).to_contain_text(text);assert student.evaluate('document.documentElement.scrollWidth<=innerWidth')
 assert not errors,errors;student_context.close();browser.close()
print('Studio/access browser passed: 10k server pagination, private rich preview, direct save/hold, actual XLSX download/openpyxl edit/upload/apply/replay rejection, frontend registration/email confirmation/pending/manual entitlement, mobile/iPad layout and real navigation.')
